"""Testes da exportação CSV/Excel (T8)."""
import csv
import io

import openpyxl

SQ = 190000000000
ND = "não disponível no dataset utilizado"


def ler_csv(conteudo: bytes):
    """Separa metadados (# chave;valor) e tabelas (## nome + cabeçalho + linhas)."""
    assert conteudo.startswith(b"\xef\xbb\xbf")        # BOM: abre certo no Excel pt-BR
    linhas = list(csv.reader(io.StringIO(conteudo.decode("utf-8-sig")), delimiter=";"))
    meta, tabelas, atual = {}, {}, None
    for l in linhas:
        if not l:
            continue
        if l[0].startswith("## "):
            atual = l[0][3:]
            tabelas[atual] = []
        elif l[0].startswith("# ") and atual is None:
            meta[l[0][2:]] = l[1]
        else:
            tabelas[atual].append(l)
    return meta, tabelas


def ler_xlsx(conteudo: bytes):
    return openpyxl.load_workbook(io.BytesIO(conteudo))


def test_busca_csv(cliente):
    r = cliente.get("/api/exportar/candidatos", params={"formato": "csv", "q": "fulano", "uf": "RJ"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="eleicoes2026_candidatos_' in r.headers["content-disposition"]
    meta, tabelas = ler_csv(r.content)
    for chave in ("fonte", "geracao_tse", "universo", "filtros_aplicados", "data_da_consulta", "natureza"):
        assert meta[chave], chave
    assert meta["geracao_tse"] == "26/09/2026 08:31:16"
    assert "q = fulano" in meta["filtros_aplicados"] and "uf = RJ" in meta["filtros_aplicados"]
    cab, *dados = tabelas["candidatos"]
    total = cliente.get("/api/candidatos", params={"q": "fulano", "uf": "RJ"}).json()["total"]
    assert len(dados) == total == int(meta["total_de_registros"])   # sem paginação
    assert "pessoa_id" not in r.content.decode("utf-8-sig")
    linha = {c: v for c, v in zip(cab, next(d for d in dados if d[0] == str(SQ)))}
    assert linha["total_bens"] == "351200,50"        # vírgula decimal
    assert linha["nm_social"] == ND                   # ausência por extenso, não zero/vazio
    assert linha["outros_registros_2026"] == f"{SQ + 1}, {SQ + 2}"
    sem_bens = {c: v for c, v in zip(cab, next(d for d in dados if d[0] == str(SQ + 9)))}
    assert sem_bens["total_bens"] == ND and sem_bens["qt_bens"] == ND
    assert "campos" in tabelas


def test_busca_xlsx(cliente):
    r = cliente.get("/api/exportar/candidatos", params={"formato": "xlsx", "na_urna": "false"})
    assert r.status_code == 200
    wb = ler_xlsx(r.content)
    assert wb.sheetnames == ["metadados", "candidatos", "campos"]
    meta = {a.value: b.value for a, b in wb["metadados"].iter_rows(min_row=2)}
    assert "na_urna = não" in meta["filtros_aplicados"]
    ws = wb["candidatos"]
    cab = [c.value for c in ws[1]]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 1 and linhas[0][0] == SQ + 8
    i = cab.index("total_bens")
    assert ws.cell(row=2, column=i + 1).number_format == '"R$" #,##0.00'
    assert ws.cell(row=2, column=i + 1).value == 351200.5
    campos = {row[0]: row[1] for row in wb["campos"].iter_rows(min_row=2, values_only=True)}
    assert campos["total_bens"] == "CALCULO" and campos["nm_urna"] == "DADO"


def test_comparacao(cliente):
    r = cliente.get("/api/exportar/comparar", params={"sq": f"{SQ},{SQ + 7}", "formato": "xlsx"})
    wb = ler_xlsx(r.content)
    assert wb.sheetnames == ["metadados", "comparacao", "diferencas", "verificacoes"]
    ws = wb["comparacao"]
    cab = [c.value for c in ws[3]]                    # linha 1 = nota, 2 = vazia
    assert cab == ["secao", "campo", "natureza", f"FULANO 0 ({SQ})", f"FULANO 7 ({SQ + 7})"]
    linhas = {row[1]: row for row in ws.iter_rows(min_row=4, values_only=True)}
    assert linhas["qt_candidaturas_anteriores"][3:] == (2, ND)
    meta = {a.value: b.value for a, b in wb["metadados"].iter_rows(min_row=2)}
    assert "Histórico não disponível" in meta["alertas"]
    textos = [row[3] for row in wb["diferencas"].iter_rows(min_row=4, values_only=True)]
    assert any("histórico não disponível" in t for t in textos)


def test_estatisticas(cliente):
    params = {"metrica": "idade_na_posse", "cd_cargo": 7, "cruzamentos": "genero,faixa_etaria"}
    r = cliente.get("/api/exportar/estatisticas", params={**params, "formato": "csv"})
    meta, tabelas = ler_csv(r.content)
    assert meta["universo"] == "Candidaturas a Deputado Estadual em todas as UFs — Eleições 2026 (n = 40)"
    api = cliente.get("/api/estatisticas", params=params).json()
    resumo = {l[0]: l[1] for l in tabelas["resumo idade_na_posse"][1:]}
    assert resumo["mediana"] == str(api["numerica"]["mediana"]).replace(".", ",")
    assert resumo["n_com_dado"] == "40"
    assert sum(int(l[2]) for l in tabelas["histograma"][1:]) == 40
    assert "genero x faixa_etaria" in tabelas
    x = ler_xlsx(cliente.get("/api/exportar/estatisticas", params={**params, "formato": "xlsx"}).content)
    assert "genero x faixa_etaria" in x.sheetnames


def test_formato_invalido(cliente):
    assert cliente.get("/api/exportar/candidatos", params={"formato": "pdf"}).status_code == 422
    assert cliente.get("/api/exportar/comparar", params={"sq": str(SQ)}).status_code == 422
