"""
Exportação de tabelas em CSV e Excel (T8), sempre com metadados: fonte, geração TSE, universo,
filtros aplicados e data da consulta.

- CSV: UTF-8 com BOM, separador ';' e vírgula decimal (abre direto no Excel em pt-BR). Os
  metadados vêm primeiro, em linhas iniciadas por '#'; cada tabela vem depois de uma linha '## nome'.
- Excel: aba "metadados" + uma aba por tabela; aba "campos" com natureza e fonte de cada coluna
  quando disponível.
- Ausência vira o texto "não disponível no dataset utilizado" — nunca célula com zero (regra 4).
- pessoa_id nunca é exportado (a verificação fica em `montar_tabela`).
"""
import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

NAO_DISPONIVEL = "não disponível no dataset utilizado"
PROIBIDAS = {"pessoa_id"}


@dataclass
class Tabela:
    nome: str
    colunas: list[str]
    linhas: list[list]
    moeda: set[str] = field(default_factory=set)   # colunas em R$ (formato no Excel)
    nota: str | None = None


@dataclass
class Exportacao:
    titulo: str
    metadados: list[tuple[str, str]]
    tabelas: list[Tabela]


def montar_tabela(nome: str, registros: list[dict], colunas: list[str], moeda=(), nota=None) -> Tabela:
    if PROIBIDAS & set(colunas):
        raise ValueError("pessoa_id nunca é exportado")
    return Tabela(nome, colunas, [[r.get(c) for c in colunas] for r in registros], set(moeda), nota)


def metadados_base(titulo: str, geracao_tse, fonte: str, universo: str | None,
                   filtros: dict | None, natureza: str, extras=()) -> list[tuple[str, str]]:
    itens = [
        ("titulo", titulo),
        ("natureza", natureza),
        ("fonte", fonte),
        ("geracao_tse", str(geracao_tse)),
        ("universo", universo or "—"),
        ("filtros_aplicados", _texto_filtros(filtros)),
        ("data_da_consulta", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("ausencia", f"Células com '{NAO_DISPONIVEL}' indicam dado ausente na base do TSE; "
                     "não equivalem a zero."),
        ("aviso", "Dados e cálculos a partir dos dados abertos do TSE. Não constituem "
                  "recomendação de voto nem avaliação de candidatos."),
    ]
    return itens + list(extras)


def _texto_filtros(filtros: dict | None) -> str:
    if not filtros:
        return "nenhum"
    partes = []
    for k, v in filtros.items():
        if v in (None, "", []):
            continue
        if isinstance(v, bool):
            v = "sim" if v else "não"
        partes.append(f"{k} = {', '.join(map(str, v)) if isinstance(v, list) else v}")
    return "; ".join(partes) or "nenhum"


# ---------------------------------------------------------------- valores
def _valor_csv(v):
    if v is None:
        return NAO_DISPONIVEL
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, (float, Decimal)):
        return f"{v}".replace(".", ",")
    if isinstance(v, (list, tuple)):
        return ", ".join(map(str, v))
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m/%Y")
    return v


def _valor_xlsx(v):
    if v is None:
        return NAO_DISPONIVEL
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (list, tuple, dict)):
        return ", ".join(map(str, v)) if not isinstance(v, dict) else str(v)
    return v


# ---------------------------------------------------------------- escritores
def para_csv(exp: Exportacao) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    for chave, valor in exp.metadados:
        w.writerow([f"# {chave}", valor])
    for t in exp.tabelas:
        w.writerow([])
        w.writerow([f"## {t.nome}"] + ([t.nota] if t.nota else []))
        w.writerow(t.colunas)
        for linha in t.linhas:
            w.writerow([_valor_csv(v) for v in linha])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def para_xlsx(exp: Exportacao, campos: list[dict] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "metadados"
    ws.append(["chave", "valor"])
    for chave, valor in exp.metadados:
        ws.append([chave, valor])
    _estilo(ws, [30, 110])
    usados = {"metadados"}
    for t in exp.tabelas:
        nome = _nome_aba(t.nome, usados)
        ws = wb.create_sheet(nome)
        if t.nota:
            ws.append([t.nota])
            ws.append([])
        ws.append(t.colunas)
        cab = ws.max_row
        for linha in t.linhas:
            ws.append([_valor_xlsx(v) for v in linha])
        for i, col in enumerate(t.colunas, start=1):
            if col in t.moeda:
                for (celula,) in ws.iter_rows(min_row=cab + 1, min_col=i, max_col=i):
                    if isinstance(celula.value, (int, float)):
                        celula.number_format = '"R$" #,##0.00'
        _estilo(ws, [max(12, min(50, len(c) + 4)) for c in t.colunas], linha_cabecalho=cab)
    if campos:
        ws = wb.create_sheet(_nome_aba("campos", usados))
        ws.append(["coluna", "natureza", "arquivo", "campo_tse", "observacao"])
        for c in campos:
            ws.append([c.get("coluna"), c.get("natureza"), c.get("arquivo"), c.get("campo_tse"),
                       c.get("observacao")])
        _estilo(ws, [28, 14, 44, 44, 90])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _nome_aba(nome: str, usados: set) -> str:
    base = "".join(ch for ch in nome if ch not in "[]:*?/\\")[:31] or "tabela"
    final, i = base, 2
    while final in usados:
        final = f"{base[:28]}_{i}"
        i += 1
    usados.add(final)
    return final


def _estilo(ws, larguras, linha_cabecalho=1):
    for cel in ws[linha_cabecalho]:
        cel.font = Font(bold=True)
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for linha in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for cel in linha:
            cel.alignment = Alignment(vertical="top", wrap_text=isinstance(cel.value, str)
                                      and len(cel.value) > 60)
    ws.freeze_panes = ws.cell(row=linha_cabecalho + 1, column=1)
